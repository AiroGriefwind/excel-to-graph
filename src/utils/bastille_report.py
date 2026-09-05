from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from .report_export import _set_cell_no_wrap, _set_run_font, format_report_number

# 平台識別關鍵字（不區分大小寫）
BASTILLE_PLATFORM_KEYWORD = "bastilleglobal"

# 博文表固定列；「欄目」取自 Excel 欄目列（新格式第 3 列，無欄目時 Excel 寫 "-"）
BASTILLE_COLUMNS = ["數目", "日期", "欄目", "標題", "瀏覽量"]

# Word 列寬（cm）；標題/欄目列允許換行，其餘列 noWrap
_COLUMN_WIDTHS_CM = [1.5, 2.3, 3.0, 6.7, 2.5]
_NO_WRAP_COLUMNS = {"數目", "日期", "瀏覽量"}
_RIGHT_ALIGN_COLUMNS = {"瀏覽量"}


def filter_bastille_rows(df: pd.DataFrame) -> pd.DataFrame:
    """篩選平台為 BastilleGlobal 的記錄，按日期升序（無日期排最後，同日期保持原順序）。"""
    working = df.copy() if df is not None else pd.DataFrame()
    for col in ("platform", "title", "views", "date", "section"):
        if col not in working.columns:
            working[col] = None

    platform_text = working["platform"].astype(str).str.strip().str.lower()
    mask = platform_text.str.contains(BASTILLE_PLATFORM_KEYWORD, na=False)
    subset = working.loc[mask].copy()
    if subset.empty:
        return subset

    sort_key = pd.to_datetime(subset["date"], errors="coerce")
    subset = subset.assign(_sort_date=sort_key).sort_values(
        "_sort_date", kind="stable", na_position="last"
    )
    return subset.drop(columns=["_sort_date"]).reset_index(drop=True)


def build_bastille_table(df: pd.DataFrame) -> pd.DataFrame:
    """生成 BastilleGlobal 博文表：數目順排、日期 YYYYMMDD、欄目取自 Excel（缺失留空）。"""
    subset = filter_bastille_rows(df)
    if subset.empty:
        return pd.DataFrame(columns=BASTILLE_COLUMNS)

    dates = pd.to_datetime(subset["date"], errors="coerce")
    views = pd.to_numeric(subset["views"], errors="coerce").fillna(0)
    sections = (
        subset["section"]
        .astype(str)
        .str.strip()
        .replace({"None": "", "nan": "", "NaN": ""})
        .tolist()
    )

    table = pd.DataFrame(
        {
            "數目": range(1, len(subset) + 1),
            "日期": [d.strftime("%Y%m%d") if pd.notna(d) else "" for d in dates],
            "欄目": sections,
            "標題": subset["title"].astype(str).tolist(),
            "瀏覽量": [int(round(v)) for v in views],
        },
        columns=BASTILLE_COLUMNS,
    )
    return table


def _row_display_values(row: pd.Series) -> list[str]:
    return [
        str(row["數目"]),
        str(row["日期"]),
        str(row["欄目"]),
        str(row["標題"]),
        format_report_number(row["瀏覽量"]),
    ]


def build_bastille_docx(
    table_df: pd.DataFrame,
    *,
    title: str = "巴士的報英文版博文列表",
    subtitle: str | None = None,
) -> bytes:
    """將博文表 DataFrame 寫成 .docx 二進位內容。"""
    doc = Document()

    heading = doc.add_heading(title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in heading.runs:
        _set_run_font(run, font_size=16)

    meta = subtitle or f"匯出時間：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
    meta_p = doc.add_paragraph(meta)
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in meta_p.runs:
        _set_run_font(run, font_size=10)

    table = doc.add_table(rows=1, cols=len(BASTILLE_COLUMNS))
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False

    header_cells = table.rows[0].cells
    for idx, col_name in enumerate(BASTILLE_COLUMNS):
        header_cells[idx].width = Cm(_COLUMN_WIDTHS_CM[idx])
        header_cells[idx].text = col_name
        if col_name in _NO_WRAP_COLUMNS:
            _set_cell_no_wrap(header_cells[idx])
        for paragraph in header_cells[idx].paragraphs:
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.RIGHT
                if col_name in _RIGHT_ALIGN_COLUMNS
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            for run in paragraph.runs:
                _set_run_font(run, font_size=11)
                run.bold = True

    for _, row in table_df.iterrows():
        cells = table.add_row().cells
        values = _row_display_values(row)
        for idx, value in enumerate(values):
            col_name = BASTILLE_COLUMNS[idx]
            cells[idx].width = Cm(_COLUMN_WIDTHS_CM[idx])
            cells[idx].text = value
            if col_name in _NO_WRAP_COLUMNS:
                _set_cell_no_wrap(cells[idx])
            for paragraph in cells[idx].paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.RIGHT
                    if col_name in _RIGHT_ALIGN_COLUMNS
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in paragraph.runs:
                    _set_run_font(run, font_size=11)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def default_bastille_filename() -> str:
    return f"巴士的報英文版博文列表_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
